"""Collect public/owner-supplied knowledge. Called by refresh_knowledge.py, NEVER per chat."""
from __future__ import annotations

import base64
import html
import json
import os
import re
import xml.etree.ElementTree as ET
from html.parser import HTMLParser
from pathlib import Path

import requests
from pypdf import PdfReader
from docx import Document
from rag.settings import PROJECT_ROOT, MANUAL_DIR

AGENT = "DigitalTwinKnowledgeRefresh/1.0 (personal research index)"
TIMEOUT = 15


# Extract readable text from feed HTML while discarding executable/style data.
class _Text(HTMLParser):
    # Accumulate visible text and track whether parsing is inside a skipped tag.
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.skip = 0

    # Preserve useful block boundaries and begin skipping scripts/styles.
    def handle_starttag(self, tag, attrs):
        if tag in ("style", "script"):
            self.skip += 1
        elif tag in ("p", "br", "li", "h1", "h2", "h3"):
            self.parts.append("\n")

    # End skipped regions and preserve closing block boundaries.
    def handle_endtag(self, tag):
        if tag in ("style", "script"):
            self.skip = max(0, self.skip - 1)
        elif tag in ("p", "li", "h1", "h2", "h3"):
            self.parts.append("\n")

    # Collect only text that is outside script and style elements.
    def handle_data(self, text):
        if not self.skip:
            self.parts.append(text)


# Normalize HTML or plain text into compact, index-ready content.
def clean_text(raw: str) -> str:
    if "<" in raw and ">" in raw:
        parser = _Text()
        parser.feed(raw)
        raw = " ".join(parser.parts)
    return re.sub(r"\s+", " ", html.unescape(raw)).strip()


# Create the common document shape used by every local and public collector.
def document(source: str, title: str, text: str, url: str = "") -> dict:
    return {"source": source, "title": title[:180], "url": url, "text": clean_text(text)[:200000]}


# Discover the fixed profile/resume sources plus owner-maintained manual files.
def local_documents() -> list[dict]:
    docs = []
    known = [
        (PROJECT_ROOT / "summary.txt", "Professional summary"),
        (PROJECT_ROOT / "Profile.pdf", "LinkedIn profile export PDF"),
        (PROJECT_ROOT / "resume.pdf", "Resume"),
        (PROJECT_ROOT / "resume.docx", "Resume"),
    ]
    for file, title in known:
        if file.exists():
            docs.extend(read_document(file, title))
    if MANUAL_DIR.exists():
        for file in sorted(MANUAL_DIR.rglob("*")):
            if (file.is_file() and "youtube" not in file.relative_to(MANUAL_DIR).parts
                    and file.name.lower() != "readme.md"
                    and file.suffix.lower() in (".md", ".txt", ".pdf", ".docx")):
                docs.extend(read_document(file, file.stem.replace("_", " ")))
    return docs


# Extract text from one supported local file without allowing a bad document to
# stop the rest of a knowledge refresh.
def read_document(path: Path, title: str) -> list[dict]:
    try:
        if path.suffix.lower() == ".pdf":
            pages = PdfReader(str(path)).pages
            text = "\n".join(page.extract_text() or "" for page in pages)
        elif path.suffix.lower() == ".docx":
            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            parts.extend(" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
            text = "\n".join(parts)
        else:
            text = path.read_text(encoding="utf-8")
    except Exception as error:
        print(f"[local] Cannot read {path.name}: {type(error).__name__}")
        return []
    if not text.strip():
        return []
    source = "linkedin" if path.name.lower() in ("profile.pdf", "linkedin_profile.txt") else "local"
    if not text.strip():
        return []
    return [document(source, title, text)]


# Collect recent article content exposed by the configured Medium RSS feed.
def fetch_medium(feed_url: str, max_items: int = 20) -> list[dict]:
    if not feed_url:
        return []
    response = requests.get(feed_url, timeout=TIMEOUT, headers={"User-Agent": AGENT})
    response.raise_for_status()
    root = ET.fromstring(response.content)
    items = root.findall("./channel/item")
    docs = []
    for item in items[:max_items]:
        title = item.findtext("title") or "Medium article"
        link = (item.findtext("link") or "").strip()
        content = item.findtext("{http://purl.org/rss/1.0/modules/content/}encoded")
        if not content:
            content = item.findtext("description") or ""
        if content.strip():
            docs.append(document("medium", title, content, link))
    return docs


# Call one public GitHub API path with optional authentication for rate limits.
def _github_get(path: str, raw: bool = False):
    headers = {"Accept": "application/vnd.github+json", "User-Agent": AGENT}
    token = os.getenv("GITHUB_TOKEN")  # Optional for API rate limits; never commit it.
    if token:
        headers["Authorization"] = f"Bearer {token}"
    response = requests.get("https://api.github.com" + path,
                            headers=headers, timeout=TIMEOUT)
    if response.status_code == 404:
        return None
    response.raise_for_status()
    return response.json()


# Collect public repository metadata and README text as project evidence.
def fetch_github(username: str, max_repos: int = 12) -> list[dict]:
    if not username:
        return []
    repos = _github_get(f"/users/{username}/repos?per_page=100&sort=updated") or []
    if not isinstance(repos, list):
        raise ValueError("Unexpected GitHub API response")
    docs = []
    for repo in repos:
        if repo.get("fork") or repo.get("private") or repo.get("archived"):
            continue
        full_name = repo.get("full_name", "")
        if not full_name.startswith(username + "/"):
            continue
        parts = [
            f"Project: {repo.get('name', '')}",
            f"Description: {repo.get('description') or 'Not provided'}",
            f"Main language: {repo.get('language') or 'Not specified'}",
            f"Topics: {', '.join(repo.get('topics') or [])}",
        ]
        try:
            readme = _github_get(f"/repos/{full_name}/readme")
            if readme and readme.get("content") and readme.get("encoding") == "base64":
                raw = base64.b64decode(readme["content"]).decode("utf-8", errors="replace")
                parts.append("README:\n" + raw[:35000])
        except (requests.RequestException, ValueError, KeyError) as error:
            print(f"[github] README skipped for {full_name}: {type(error).__name__}")
        docs.append(document("github", f"GitHub: {full_name}", "\n".join(parts), repo.get("html_url", "")))
        if len(docs) >= max_repos:
            break
    return docs


# Collect titles and descriptions from a public YouTube channel feed.
def fetch_youtube(channel_id: str, max_items: int = 20) -> list[dict]:
    if not channel_id:
        return []
    url = "https://www.youtube.com/feeds/videos.xml"
    response = requests.get(url, params={"channel_id": channel_id},
                            headers={"User-Agent": AGENT}, timeout=TIMEOUT)
    response.raise_for_status()
    root = ET.fromstring(response.content)
    atom = "{http://www.w3.org/2005/Atom}"
    yt = "{http://www.youtube.com/xml/schemas/2015}"
    media = "{http://search.yahoo.com/mrss/}"
    docs = []
    for entry in root.findall(atom + "entry")[:max_items]:
        video_id = entry.findtext(yt + "videoId") or ""
        title = entry.findtext(atom + "title") or "YouTube video"
        description = entry.findtext(f"{media}group/{media}description") or ""
        if not video_id:
            continue
        docs.append(document("youtube", title, f"Video title: {title}. Description: {description}",
                             f"https://www.youtube.com/watch?v={video_id}"))
    return docs


# Remove timing, markup, and repeated caption lines from VTT/SRT text.
def _subtitle_text(text: str) -> str:
    result = []
    for line in text.splitlines():
        line = line.strip().lstrip("\ufeff")
        if not line or line == "WEBVTT" or line.isdecimal() or "-->" in line:
            continue
        if re.fullmatch(r"\d\d:\d\d(?::\d\d)?[.,]\d+", line):
            continue
        if line.startswith(("NOTE ", "STYLE", "Kind:")):
            continue
        line = re.sub(r"<[^>]+>", "", line)
        if line and (not result or result[-1] != line):
            result.append(line)
    return " ".join(result)


# Convert owner-supplied caption files into documents with verified video URLs.
def transcript_documents() -> list[dict]:
    directory = MANUAL_DIR / "youtube"
    if not directory.exists():
        return []
    docs = []
    manifest_file = directory / "videos.json"
    titles = {}
    if manifest_file.exists():
        for video in json.loads(manifest_file.read_text(encoding="utf-8")):
            titles[video["id"]] = video.get("title", "")
    for file in sorted(directory.iterdir()):
        if file.name.lower() in ("readme.md", "videos.json"):
            continue
        if file.suffix.lower() not in (".txt", ".md", ".vtt", ".srt"):
            continue
        video_id = file.stem
        if not re.fullmatch(r"[A-Za-z0-9_-]{11}", video_id):
            print(f"[youtube] Skipped {file.name}: filename must equal 11-char video ID")
            continue
        text = _subtitle_text(file.read_text(encoding="utf-8")) if file.suffix.lower() in (".vtt", ".srt") else file.read_text(encoding="utf-8")
        title = titles.get(video_id) or f"Video transcript {video_id}"
        docs.append(document("youtube_transcript", title, text,
                             f"https://www.youtube.com/watch?v={video_id}"))
    return docs
